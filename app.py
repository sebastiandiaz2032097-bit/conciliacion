import streamlit as st
import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Portal de Conciliación Contable", layout="wide")

st.title("📊 Portal de Conciliación: Balance General vs. Amortización")
st.write("Carga tus archivos `.txt`, mapea tus columnas, aplica filtros por cuenta y descarga el reporte detallado en Excel.")

# Sidebar - Configuración
st.sidebar.header("⚙️ Configuración Global")
separador_opcion = st.sidebar.selectbox(
    "Separador de archivos TXT:",
    ["Punto y coma (;)", "Coma (,)", "Tabulación (Tab)"]
)

mapa_separadores = {
    "Punto y coma (;)": ";",
    "Coma (,)": ",",
    "Tabulación (Tab)": "\t"
}
sep = mapa_separadores[separador_opcion]

# Carga de archivos
col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Balance General (.txt)")
    file_bg = st.file_uploader("Cargar Balance", type=["txt"], key="bg")

with col2:
    st.subheader("2. Auxiliar de Amortización (.txt)")
    file_amort = st.file_uploader("Cargar Amortización", type=["txt"], key="amort")

def limpiar_monto(valor_str):
    if not valor_str:
        return 0.0
    val = str(valor_str).strip().replace("$", "").replace(" ", "")
    if "." in val and "," in val:
        val = val.replace(".", "").replace(",", ".")
    elif "," in val:
        val = val.replace(",", ".")
    try:
        return float(val)
    except ValueError:
        return 0.0

if file_bg and file_amort:
    content_bg = file_bg.getvalue().decode('utf-8-sig', errors='ignore')
    content_amort = file_amort.getvalue().decode('utf-8-sig', errors='ignore')
    
    rows_bg = list(csv.reader(io.StringIO(content_bg), delimiter=sep))
    rows_amort = list(csv.reader(io.StringIO(content_amort), delimiter=sep))

    if len(rows_bg) > 1 and len(rows_amort) > 1:
        headers_bg = [f"Col {i}: {col}" for i, col in enumerate(rows_bg[0])]
        headers_amort = [f"Col {i}: {col}" for i, col in enumerate(rows_amort[0])]

        st.write("---")
        st.subheader("⚙️ Mapeo de Columnas (Requeridas y Opcionales)")

        m1, m2 = st.columns(2)
        with m1:
            st.write("**Balance General**")
            col_cedula_bg = st.selectbox("Cédula/NIT (Balance):", range(len(rows_bg[0])), format_func=lambda x: headers_bg[x], index=0)
            col_valor_bg = st.selectbox("Valor/Saldo (Balance):", range(len(rows_bg[0])), format_func=lambda x: headers_bg[x], index=min(1, len(rows_bg[0])-1))
            
            # Columnas opcionales
            opc_cuenta_bg = st.checkbox("Incluir columna Cuenta Contable (Balance)", value=False, key="cb_cuenta_bg")
            col_cuenta_bg = st.selectbox("Cuenta Contable (Balance):", range(len(rows_bg[0])), format_func=lambda x: headers_bg[x], index=0) if opc_cuenta_bg else None

            opc_nombre_bg = st.checkbox("Incluir columna Nombre / Tercero (Balance)", value=False, key="cb_nombre_bg")
            col_nombre_bg = st.selectbox("Nombre / Tercero (Balance):", range(len(rows_bg[0])), format_func=lambda x: headers_bg[x], index=0) if opc_nombre_bg else None

        with m2:
            st.write("**Auxiliar de Amortización**")
            col_cedula_amort = st.selectbox("Cédula/NIT (Amortización):", range(len(rows_amort[0])), format_func=lambda x: headers_amort[x], index=0)
            col_valor_amort = st.selectbox("Valor/Saldo (Amortización):", range(len(rows_amort[0])), format_func=lambda x: headers_amort[x], index=min(1, len(rows_amort[0])-1))
            
            # Columnas opcionales
            opc_cuenta_amort = st.checkbox("Incluir columna Cuenta Contable (Amortización)", value=False, key="cb_cuenta_amort")
            col_cuenta_amort = st.selectbox("Cuenta Contable (Amortización):", range(len(rows_amort[0])), format_func=lambda x: headers_amort[x], index=0) if opc_cuenta_amort else None

            opc_nombre_amort = st.checkbox("Incluir columna Nombre / Tercero (Amortización)", value=False, key="cb_nombre_amort")
            col_nombre_amort = st.selectbox("Nombre / Tercero (Amortización):", range(len(rows_amort[0])), format_func=lambda x: headers_amort[x], index=0) if opc_nombre_amort else None

        # --- MEJORA 1: Filtro opcional por Cuenta Contable ---
        st.write("---")
        st.subheader("🔍 Filtro por Cuenta Contable (Opcional)")
        
        # Extraer lista de cuentas para el filtro
        cuentas_disponibles = set()
        if col_cuenta_bg is not None:
            for r in rows_bg[1:]:
                if len(r) > col_cuenta_bg and r[col_cuenta_bg].strip():
                    cuentas_disponibles.add(r[col_cuenta_bg].strip())
        if col_cuenta_amort is not None:
            for r in rows_amort[1:]:
                if len(r) > col_cuenta_amort and r[col_cuenta_amort].strip():
                    cuentas_disponibles.add(r[col_cuenta_amort].strip())

        filtro_cuentas = []
        if cuentas_disponibles:
            filtro_cuentas = st.multiselect("Selecciona una o varias cuentas para filtrar la conciliación (deja vacío para procesar todas):", sorted(list(cuentas_disponibles)))
        else:
            st.caption("💡 Activa el mapeo de 'Cuenta Contable' en los archivos arriba para habilitar el filtro por cuenta.")

        if st.button("🔍 Realizar Conciliación", type="primary"):
            saldos_bg = {}
            nombres_terceros = {}
            cuentas_terceros = {}

            # Procesar Balance General
            for row in rows_bg[1:]:
                if len(row) > max(col_cedula_bg, col_valor_bg):
                    cedula = str(row[col_cedula_bg]).strip()
                    monto = limpiar_monto(row[col_valor_bg])
                    cuenta = str(row[col_cuenta_bg]).strip() if col_cuenta_bg is not None and len(row) > col_cuenta_bg else ""
                    nombre = str(row[col_nombre_bg]).strip() if col_nombre_bg is not None and len(row) > col_nombre_bg else ""

                    # Aplicar filtro si existe
                    if filtro_cuentas and cuenta and cuenta not in filtro_cuentas:
                        continue

                    if cedula:
                        saldos_bg[cedula] = saldos_bg.get(cedula, 0.0) + monto
                        if nombre and cedula not in nombres_terceros:
                            nombres_terceros[cedula] = nombre
                        if cuenta and cedula not in cuentas_terceros:
                            cuentas_terceros[cedula] = cuenta

            # Procesar Amortización
            saldos_amort = {}
            for row in rows_amort[1:]:
                if len(row) > max(col_cedula_amort, col_valor_amort):
                    cedula = str(row[col_cedula_amort]).strip()
                    monto = limpiar_monto(row[col_valor_amort])
                    cuenta = str(row[col_cuenta_amort]).strip() if col_cuenta_amort is not None and len(row) > col_cuenta_amort else ""
                    nombre = str(row[col_nombre_amort]).strip() if col_nombre_amort is not None and len(row) > col_nombre_amort else ""

                    # Aplicar filtro si existe
                    if filtro_cuentas and cuenta and cuenta not in filtro_cuentas:
                        continue

                    if cedula:
                        saldos_amort[cedula] = saldos_amort.get(cedula, 0.0) + monto
                        if nombre and cedula not in nombres_terceros:
                            nombres_terceros[cedula] = nombre
                        if cuenta and cedula not in cuentas_terceros:
                            cuentas_terceros[cedula] = cuenta

            todas_cedulas = set(saldos_bg.keys()).union(set(saldos_amort.keys()))

            diferencias = []
            total_bg_sum = 0.0
            total_amort_sum = 0.0

            for ced in todas_cedulas:
                val_bg = saldos_bg.get(ced, 0.0)
                val_amort = saldos_amort.get(ced, 0.0)
                diferencia = val_bg - val_amort

                total_bg_sum += val_bg
                total_amort_sum += val_amort

                if abs(diferencia) > 0.01:
                    diferencias.append({
                        "Cédula / NIT": ced,
                        "Nombre / Tercero": nombres_terceros.get(ced, "N/A"),
                        "Cuenta Contable": cuentas_terceros.get(ced, "N/A"),
                        "Saldo Balance": val_bg,
                        "Saldo Amortización": val_amort,
                        "Diferencia": diferencia
                    })

            # --- MEJORA 3: Resumen Ejecutivo en Pantalla ---
            st.write("---")
            st.subheader("📈 Resumen Ejecutivo de la Conciliación")
            
            kpi1, kpi2, kpi3, kpi4 = st.columns(4)
            kpi1.metric("Total Balance General", f"$ {total_bg_sum:,.2f}")
            kpi2.metric("Total Amortización", f"$ {total_amort_sum:,.2f}")
            kpi3.metric("Diferencia Neta Global", f"$ {(total_bg_sum - total_amort_sum):,.2f}")
            kpi4.metric("Terceros Descuadrados", len(diferencias))

            if diferencias:
                st.warning(f"Se identificaron **{len(diferencias)}** terceros con diferencias numéricas.")

                # --- MEJORA 2: Generar Excel con Formato Profesional (openpyxl) ---
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Diferencias Conciliación"

                # Estilos visuales
                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Azul oscuro
                total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")  # Azul claro
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                bold_font = Font(name="Calibri", size=11, bold=True)
                regular_font = Font(name="Calibri", size=11)
                
                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'),
                    right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'),
                    bottom=Side(style='thin', color='D9D9D9')
                )

                # Encabezados en Excel (Con Mejora 4: Columnas Adicionales)
                headers = ["Cédula / NIT", "Nombre / Tercero", "Cuenta Contable", "Saldo Balance", "Saldo Amortización", "Diferencia"]
                ws.append(headers)

                for col_num, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Insertar filas de datos
                for item in diferencias:
                    ws.append([
                        item["Cédula / NIT"],
                        item["Nombre / Tercero"],
                        item["Cuenta Contable"],
                        item["Saldo Balance"],
                        item["Saldo Amortización"],
                        item["Diferencia"]
                    ])

                # Aplicar formatos numéricos, bordes y alineaciones
                row_start = 2
                row_end = len(diferencias) + 1

                for row in range(row_start, row_end + 1):
                    ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
                    ws.cell(row=row, column=2).font = regular_font
                    ws.cell(row=row, column=3).alignment = Alignment(horizontal="center")

                    # Formato de moneda contable para los números
                    for col in range(4, 7):
                        c = ws.cell(row=row, column=col)
                        c.number_format = '"$"#,##0.00;[Red]("$"#,##0.00);"-"'
                        c.font = regular_font
                        c.border = thin_border

                # Fila de Totales al final
                total_row = row_end + 1
                ws.cell(row=total_row, column=1, value="TOTAL GENERAL").font = bold_font
                ws.cell(row=total_row, column=1).alignment = Alignment(horizontal="center")
                
                ws.cell(row=total_row, column=4, value=f"=SUM(D2:D{row_end})")
                ws.cell(row=total_row, column=5, value=f"=SUM(E2:E{row_end})")
                ws.cell(row=total_row, column=6, value=f"=SUM(F2:F{row_end})")

                for col in range(1, 7):
                    c = ws.cell(row=total_row, column=col)
                    c.fill = total_fill
                    c.font = bold_font
                    c.border = thin_border
                    if col >= 4:
                        c.number_format = '"$"#,##0.00;[Red]("$"#,##0.00);"-"'

                # Autoajustar ancho de columnas en Excel
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

                # Guardar Excel en buffer para descargar
                excel_buffer = io.BytesIO()
                wb.save(excel_buffer)
                excel_buffer.seek(0)

                # Vista previa en Streamlit
                st.write("**Vista previa de las diferencias encontradas:**")
                st.dataframe(diferencias, use_container_width=True)

                # Botón de descarga
                st.download_button(
                    label="📥 Descargar Reporte Profesional en Excel (.xlsx)",
                    data=excel_buffer,
                    file_name="Reporte_Conciliacion_Diferencias.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            else:
                st.success("🎉 ¡Excelente! No se encontraron diferencias entre el Balance y la Amortización para los criterios seleccionados.")